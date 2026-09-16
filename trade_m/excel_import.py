from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from io import BytesIO
from zipfile import BadZipFile, LargeZipFile, ZipFile

from openpyxl import load_workbook


MAX_IMPORT_ROWS = 200
MAX_UNCOMPRESSED_BYTES = 20 * 1024 * 1024


@dataclass(frozen=True)
class ImportedStock:
    row_number: int
    tradingsymbol: str
    percentage: Decimal
    exchange: str


@dataclass(frozen=True)
class ParsedStockWorkbook:
    stocks: list[ImportedStock]
    skipped: int


class StockWorkbookError(ValueError):
    def __init__(self, errors: list[str]) -> None:
        self.errors = errors
        super().__init__(" ".join(errors))


def _header_key(value: object) -> str:
    return "".join(character for character in str(value or "").lower() if character.isalnum())


HEADER_ALIASES = {
    "tradingsymbol": {"tradingsymbol", "symbol", "stocksymbol", "stockname", "stock"},
    "percentage": {"percentage", "percent", "pct", "dailypercentage", "movepercentage"},
    "exchange": {"exchange", "market"},
    "enabled": {"enabled", "active", "monitor", "include"},
}


def _find_headers(sheet: object) -> tuple[int, dict[str, int]]:
    max_header_row = min(sheet.max_row or 10, 10)
    for row_number, row in enumerate(
        sheet.iter_rows(min_row=1, max_row=max_header_row), start=1
    ):
        found: dict[str, int] = {}
        for index, cell in enumerate(row):
            key = _header_key(cell.value)
            for canonical, aliases in HEADER_ALIASES.items():
                if key in aliases and canonical not in found:
                    found[canonical] = index
        if "tradingsymbol" in found and "percentage" in found:
            return row_number, found
    raise StockWorkbookError(
        ["The workbook must contain Trading Symbol and Percentage columns."]
    )


def _enabled(value: object, row_number: int) -> bool:
    if value is None or str(value).strip() == "":
        return True
    if isinstance(value, bool):
        return value
    normalized = str(value).strip().lower()
    if normalized in {"true", "yes", "y", "1", "enabled", "active"}:
        return True
    if normalized in {"false", "no", "n", "0", "disabled", "inactive"}:
        return False
    raise ValueError(f"Row {row_number}: Enabled must be TRUE or FALSE.")


def _percentage(value: object, number_format: str, row_number: int) -> Decimal:
    if value is None or str(value).strip() == "":
        raise ValueError(f"Row {row_number}: Percentage is required.")
    if isinstance(value, bool):
        raise ValueError(f"Row {row_number}: Percentage must be a number.")
    try:
        result = Decimal(str(value).strip())
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"Row {row_number}: Percentage must be a number.") from exc
    if "%" in (number_format or ""):
        result *= Decimal("100")
    if result <= 0 or result > 50:
        raise ValueError(
            f"Row {row_number}: Percentage must be greater than 0 and at most 50."
        )
    return result


def parse_stock_workbook(content: bytes) -> ParsedStockWorkbook:
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
                raise StockWorkbookError(
                    ["The expanded Excel workbook is too large to import safely."]
                )
    except StockWorkbookError:
        raise
    except (BadZipFile, LargeZipFile, OSError, ValueError) as exc:
        raise StockWorkbookError(
            ["The uploaded file is not a readable .xlsx workbook."]
        ) from exc

    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except Exception as exc:
        raise StockWorkbookError(
            ["The uploaded file is not a readable .xlsx workbook."]
        ) from exc

    try:
        if not workbook.worksheets:
            raise StockWorkbookError(["The workbook does not contain a worksheet."])
        sheet = workbook["Stocks"] if "Stocks" in workbook.sheetnames else workbook.active
        header_row, headers = _find_headers(sheet)
        stocks: list[ImportedStock] = []
        errors: list[str] = []
        skipped = 0
        seen: set[tuple[str, str]] = set()
        populated_rows = 0

        for row_number, row in enumerate(
            sheet.iter_rows(min_row=header_row + 1), start=header_row + 1
        ):
            def cell(name: str):
                index = headers.get(name)
                return row[index] if index is not None and index < len(row) else None

            symbol_cell = cell("tradingsymbol")
            percentage_cell = cell("percentage")
            exchange_cell = cell("exchange")
            enabled_cell = cell("enabled")
            symbol_value = symbol_cell.value if symbol_cell is not None else None
            percentage_value = percentage_cell.value if percentage_cell is not None else None

            if (symbol_value is None or str(symbol_value).strip() == "") and (
                percentage_value is None or str(percentage_value).strip() == ""
            ):
                continue

            populated_rows += 1
            if populated_rows > MAX_IMPORT_ROWS:
                errors.append(f"The workbook can contain at most {MAX_IMPORT_ROWS} stock rows.")
                break

            try:
                enabled = _enabled(
                    enabled_cell.value if enabled_cell is not None else None, row_number
                )
                if not enabled:
                    skipped += 1
                    continue

                symbol = str(symbol_value or "").strip().upper()
                if not symbol:
                    raise ValueError(f"Row {row_number}: Trading Symbol is required.")
                if symbol.startswith("="):
                    raise ValueError(
                        f"Row {row_number}: Trading Symbol must be plain text, not a formula."
                    )
                if len(symbol) > 40:
                    raise ValueError(
                        f"Row {row_number}: Trading Symbol must be 40 characters or fewer."
                    )

                raw_exchange = exchange_cell.value if exchange_cell is not None else None
                exchange = str(raw_exchange or "NSE").strip().upper() or "NSE"
                if exchange not in {"NSE", "BSE"}:
                    raise ValueError(f"Row {row_number}: Exchange must be NSE or BSE.")

                percentage = _percentage(
                    percentage_value,
                    percentage_cell.number_format if percentage_cell is not None else "",
                    row_number,
                )
                duplicate_key = (exchange, symbol)
                if duplicate_key in seen:
                    raise ValueError(
                        f"Row {row_number}: {exchange}:{symbol} appears more than once."
                    )
                seen.add(duplicate_key)
                stocks.append(
                    ImportedStock(
                        row_number=row_number,
                        tradingsymbol=symbol,
                        percentage=percentage,
                        exchange=exchange,
                    )
                )
            except ValueError as exc:
                errors.append(str(exc))

        if errors:
            raise StockWorkbookError(errors)
        if not stocks:
            raise StockWorkbookError(
                ["The workbook contains no enabled stock rows to import."]
            )
        return ParsedStockWorkbook(stocks=stocks, skipped=skipped)
    finally:
        workbook.close()
